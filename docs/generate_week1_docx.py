# -*- coding: utf-8 -*-
"""生成《第一周进度汇报》Word 文档(docx),自动嵌入 outputs/week1 下的图表。

用法:
    python generate_week1_docx.py
"""
import os

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))          # docs/
PROJECT = os.path.dirname(BASE)                            # 项目根目录
OUT = os.path.join(PROJECT, "outputs", "week1")
DOCX = os.path.join(BASE, "第1周进度汇报.docx")

CN_FONT = "微软雅黑"

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = CN_FONT
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)


def set_cn(run, size=11, bold=False, font=CN_FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def heading(text, level=1):
    h = doc.add_heading("", level=level)
    r = h.add_run(text)
    set_cn(r, size={1: 16, 2: 13, 3: 12}[level], bold=True)
    r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
    return h


def para(text, bold=False, size=11):
    p = doc.add_paragraph()
    set_cn(p.add_run(text), size=size, bold=bold)
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    set_cn(p.add_run(text))
    return p


def code(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Consolas"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    r.font.size = Pt(9)
    return p


def image(path, caption=""):
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(5.3))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if caption:
            para(caption, size=9).alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        para(f"[缺少图片: {os.path.basename(path)}]", size=9)


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                set_cn(r, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    return t


# ================= 标题 =================
tp = doc.add_paragraph()
tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_cn(tp.add_run("手写文字识别系统(HTR)—— 第一周进度汇报"), size=18, bold=True)

sp = doc.add_paragraph()
sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_cn(sp.add_run("汇报人:张加豪    日期:2026-09-10    项目:基于深度学习的手写文字识别系统"), size=10)

# ================= 一、概述 =================
heading("一、本周工作概述", 1)
para("本周完成了项目的环境搭建与第一个里程碑任务 —— MNIST 手写数字识别,"
     "跑通了 PyTorch 从数据加载、模型构建、训练、验证到测试评估的完整流程。")
para("核心成果:测试集准确率 99.16%,达到并超过验收标准(≥98%)。", bold=True)
for t in ["深度学习环境配置(GPU 版 PyTorch)",
          "项目结构搭建 + GitHub 仓库 + 首次提交",
          "MNIST 数据集加载与预处理",
          "CNN 模型实现与训练(5 个 epoch)",
          "损失/准确率曲线、混淆矩阵、错误样本分析"]:
    bullet(t)

# ================= 二、环境 =================
heading("二、环境与工程搭建", 1)
heading("2.1 软硬件环境", 2)
table(["项目", "配置"], [
    ["操作系统", "Windows 11"],
    ["GPU", "NVIDIA GeForce RTX 3050 Laptop(4GB 显存)"],
    ["Python", "3.13.9(conda 环境 tuduipytorch)"],
    ["深度学习框架", "PyTorch 2.9.1 + CUDA 13.0"],
    ["其他依赖", "torchvision 0.24.1、matplotlib、scikit-learn 等"],
])
heading("2.2 项目结构", 2)
code("Handwritten-Text-Recognition/\n"
     "├── week1_mnist/            # 第1周:MNIST 数字识别(已完成)\n"
     "├── week2_emnist/           # 第2周:EMNIST 字符识别(待做)\n"
     "├── week3_iam_crnn/         # 第3周:IAM 单词级识别(待做)\n"
     "├── week4_transformer_ocr/  # 第4周:Transformer OCR 调研(待做)\n"
     "├── docs/                   # 文档与报告\n"
     "├── outputs/                # 训练产物(模型权重、图表)\n"
     "└── data/                   # 数据集(不入库)")
heading("2.3 代码管理", 2)
bullet("已建立 GitHub 仓库 zjh-eng/Practice,使用 main 分支")
bullet("完成首次提交(12 个文件),后续将保持持续提交记录")

# ================= 三、数据集 =================
heading("三、数据集:MNIST", 1)
for t in ["内容:0~9 共 10 类手写数字",
          "规模:60,000 张训练图 + 10,000 张测试图",
          "尺寸:28×28 单通道灰度图",
          "预处理:ToTensor(归一化到 [0,1]) + Normalize(均值 0.1307、标准差 0.3081)",
          "划分:训练集按 9:1 切出验证集,测试集独立用于最终评估"]:
    bullet(t)

# ================= 四、模型 =================
heading("四、模型结构:CNN", 1)
para("采用「两层卷积 + 两层全连接」的经典 CNN,总参数量 421,642(约 0.42M)。")
code("Input [1, 28, 28]\n"
     "  → Conv(1→32, 3×3) → ReLU → MaxPool(2)     # 28×28 → 14×14\n"
     "  → Conv(32→64, 3×3) → ReLU → MaxPool(2)    # 14×14 → 7×7\n"
     "  → Flatten → FC(3136→128) → ReLU → Dropout(0.5)\n"
     "  → FC(128→10)")
para("设计要点:卷积层负责提取局部特征,池化下采样压缩维度,全连接层完成分类;"
     "Dropout 用于抑制过拟合。")

# ================= 五、训练 =================
heading("五、训练配置与过程", 1)
table(["超参数", "取值"], [
    ["优化器", "Adam"],
    ["学习率", "1e-3"],
    ["Batch Size", "64"],
    ["训练轮数", "5"],
    ["损失函数", "CrossEntropyLoss"],
    ["随机种子", "42"],
])
para("训练过程中损失逐步下降、准确率稳步上升,5 个 epoch 后验证集准确率收敛至 99% 左右。")
image(os.path.join(OUT, "training_curves.png"), "图1 训练/验证的损失与准确率曲线")

# ================= 六、结果 =================
heading("六、实验结果", 1)
para("测试集准确率:99.16%(10,000 张测试图,共 84 张识别错误)。")
image(os.path.join(OUT, "confusion_matrix.png"), "图2 测试集混淆矩阵")

# ================= 七、错误分析 =================
heading("七、错误样本分析", 1)
para("对 84 个错误样本做了归类统计,发现错误主要集中在字形相近的数字之间:")
table(["排名", "混淆对(真值 → 预测)", "次数"], [
    ["1", "4 → 9", "11"],
    ["2", "3 → 5", "8"],
    ["3", "2 → 7", "6"],
    ["4", "7 → 2", "5"],
    ["5", "9 → 5", "4"],
])
para("分析结论:4 与 9、3 与 5、2 与 7 这几组数字在手写体中笔画形态容易变形、高度相似,"
     "是 CNN 出错的主要方向。")
image(os.path.join(OUT, "error_analysis.png"), "图3 错误样本可视化(红字标注真值/预测值)")

# ================= 八、AI 辅助 =================
heading("八、AI 辅助开发情况", 1)
para("按照任务书要求,本周全程使用 Claude Code 辅助开发,主要完成了:")
for t in ["项目结构设计与 README、依赖清单的编写",
          "CNN 模型与训练/评估脚本的代码编写",
          "环境问题的诊断(Windows OpenMP 冲突、conda 环境踩坑等)"]:
    bullet(t)
para("AI 辅助代码占比估计:约 70%~80%(由 AI 生成后经人工理解、审查与调整)。")

# ================= 九、问题 =================
heading("九、遇到的问题与解决", 1)
table(["问题", "原因", "解决方案"], [
    ["导入 torch 报 libiomp5md.dll 冲突", "numpy(MKL)与 torch 的 OpenMP 运行库重复",
     "脚本顶部设置 KMP_DUPLICATE_LIB_OK=TRUE"],
    ["PyCharm 运行路径错乱", "脚本用相对路径,PyCharm 工作目录默认为项目根",
     "改为基于 __file__ 推导项目根目录"],
    ["conda 新环境装 PyTorch 失败", "Windows 260 字符路径上限导致解压失败",
     "复用已有 GPU 环境 tuduipytorch"],
])

# ================= 十、下周计划 =================
heading("十、下周计划", 1)
for i, t in enumerate([
    "EMNIST 字符识别:将任务从纯数字扩展到字母+数字混合识别 —— 实现 MLP/CNN/ResNet18 对比、"
    "引入数据增强、产出准确率/训练时间/参数量对比表",
    "提前申请 IAM 数据集(第3周使用,审批周期较长)",
    "保持代码的持续提交与文档同步更新",
], start=1):
    p = doc.add_paragraph(style="List Number")
    set_cn(p.add_run(t))

# ================= 附:复现 =================
heading("附:复现方式", 1)
code("conda activate tuduipytorch\n"
     "cd week1_mnist\n"
     "python train.py      # 训练 + 保存模型 + 画曲线\n"
     "python evaluate.py   # 测试评估 + 混淆矩阵 + 错样本分析")
para("产物位置:模型权重与三张分析图均在 outputs/week1/ 目录下。", size=10)

doc.save(DOCX)
print(f"已生成: {DOCX}")
